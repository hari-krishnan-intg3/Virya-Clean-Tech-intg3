import csv
import json
import time
import urllib.error
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


FOLDER = Path(__file__).resolve().parent
CONFIG_PATH = FOLDER / "madeinusa_filters.json"
OUTPUT_STEM = "madeinusa_extracted"

EMPTY_QUERY = {
    "searchText": "",
    "productSelected": False,
    "orderBy": "",
    "ip": "",
    "country": "",
    "currentlyExporting": "",
    "productionLocations": [],
    "businessTypes": [],
    "tradePreferencePrograms": [],
    "companySizes": [],
    "exportingCountries": [],
    "exportingRegions": [],
    "products": [],
}


def fetch_page(api_url: str, query: dict[str, Any]) -> dict[str, Any]:
    request = urllib.request.Request(
        api_url,
        data=json.dumps(query).encode("utf-8"),
        headers={
            "Accept": "application/json",
            "Content-Type": "application/json",
            "User-Agent": "MadeInUSADirectoryResearch/1.0",
        },
        method="POST",
    )

    for attempt in range(3):
        try:
            with urllib.request.urlopen(request, timeout=30) as response:
                return json.load(response)
        except (urllib.error.URLError, TimeoutError) as error:
            if attempt == 2:
                raise RuntimeError(f"Directory request failed: {error}") from error
            time.sleep(1.5 * (attempt + 1))

    raise RuntimeError("Directory request failed after retries")


def json_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def normalize_record(
    result: dict[str, Any],
    state: dict[str, Any],
    filter_config: dict[str, Any],
    page_number: int,
    extracted_at: str,
    source_url: str,
    api_url: str,
) -> dict[str, Any]:
    producer = result.get("usProducer") or {}
    contact = result.get("contactInformation") or {}
    products = result.get("products") or []
    query_values = {
        **EMPTY_QUERY,
        "pageSize": filter_config["page_size"],
        "page": page_number,
        "hqstates": state["state_abbreviation"],
    }

    return {
        "record_id": result.get("id", ""),
        "state": state["state_name"],
        "state_abbreviation": state["state_abbreviation"],
        "filter_id": state["filter_id"],
        "filter_applied": json.dumps(state["filter_applied"], ensure_ascii=False),
        "company_name": producer.get("companyName") or contact.get("organizationName", ""),
        "contact_first_name": contact.get("firstName", ""),
        "contact_last_name": contact.get("lastName", ""),
        "contact_title": contact.get("title", ""),
        "email": contact.get("email", ""),
        "phone": contact.get("phone", ""),
        "fax": contact.get("fax", ""),
        "website": contact.get("website", ""),
        "address": contact.get("address1", ""),
        "address_2": contact.get("address2", ""),
        "city": contact.get("city", ""),
        "contact_state": contact.get("state", ""),
        "postal_code": contact.get("postalCode", ""),
        "country": contact.get("country", ""),
        "company_size": producer.get("companySize", ""),
        "business_types": json_value(producer.get("organizationType", [])),
        "currently_exporting": producer.get("exporting", ""),
        "export_percentage": json_value(producer.get("exportPercentage", [])),
        "imported_materials": producer.get("importedMaterials", ""),
        "minimum_starting_order": producer.get("minimumStartingOrder", ""),
        "minimum_order_unit": producer.get("listValues", ""),
        "turnaround_days": producer.get("turnAroundDays", ""),
        "products": json_value(products),
        "export_destinations": json_value(result.get("exportDestinations", [])),
        "export_regions": json_value(result.get("exportRegions", [])),
        "trade_preference_programs": json_value(result.get("preferencePrograms", [])),
        "source_url": source_url,
        "source_api_url": api_url,
        "source_record_id": result.get("id", ""),
        "source_query": json.dumps(query_values, ensure_ascii=False),
        "extraction_timestamp_utc": extracted_at,
        "source_page": page_number,
    }


def extract(filter_config: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    records: list[dict[str, Any]] = []
    summaries: list[dict[str, Any]] = []
    extracted_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    target = filter_config["target_records_per_state"]

    for state in filter_config["filters"]:
        page_number = 1
        available = 0
        state_records: list[dict[str, Any]] = []

        while len(state_records) < target:
            query = {
                **EMPTY_QUERY,
                "pageSize": filter_config["page_size"],
                "page": page_number,
                "hqstates": state["state_abbreviation"],
            }
            response = fetch_page(filter_config["search_api_url"], query)
            available = int(response.get("resultsTotal", 0))
            results = response.get("results") or []
            if not results:
                break

            remaining = target - len(state_records)
            for result in results[:remaining]:
                state_records.append(
                    normalize_record(
                        result,
                        state,
                        filter_config,
                        page_number,
                        extracted_at,
                        filter_config["source_url"],
                        filter_config["search_api_url"],
                    )
                )

            print(
                f"{state['filter_id']}: page {page_number}, "
                f"{min(len(results), remaining)} records ({len(state_records)}/{target})"
            )
            page_number += 1
            if len(state_records) >= available:
                break
            time.sleep(filter_config["request_delay_seconds"])

        records.extend(state_records)
        summaries.append(
            {
                "state": state["state_name"],
                "state_abbreviation": state["state_abbreviation"],
                "filter_id": state["filter_id"],
                "filter_applied": json.dumps(state["filter_applied"], ensure_ascii=False),
                "target_records": target,
                "available_at_source": available,
                "records_extracted": len(state_records),
                "status": "Complete" if len(state_records) >= target else "Partial - source has fewer records",
                "extraction_timestamp_utc": extracted_at,
                "source_url": filter_config["source_url"],
            }
        )

    return records, summaries


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    with path.open("w", newline="", encoding="utf-8-sig") as output:
        writer = csv.DictWriter(output, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def write_workbook(path: Path, records: list[dict[str, Any]], summaries: list[dict[str, Any]]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise RuntimeError("Excel output requires openpyxl. Install requirements_madeinusa.txt.") from error

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_rows = summaries + [
        {
            "state": "TOTAL",
            "state_abbreviation": "",
            "filter_id": "ALL",
            "filter_applied": "4 state filters",
            "target_records": sum(row["target_records"] for row in summaries),
            "available_at_source": sum(row["available_at_source"] for row in summaries),
            "records_extracted": sum(row["records_extracted"] for row in summaries),
            "status": "Complete" if all(row["status"] == "Complete" for row in summaries) else "Partial - source limits",
            "extraction_timestamp_utc": summaries[0]["extraction_timestamp_utc"] if summaries else "",
            "source_url": summaries[0]["source_url"] if summaries else "",
        }
    ]
    append_table(summary_sheet, summary_rows)

    all_records_sheet = workbook.create_sheet("All Records")
    append_table(all_records_sheet, records)

    for summary in summaries:
        sheet = workbook.create_sheet(summary["filter_id"])
        state_records = [row for row in records if row["filter_id"] == summary["filter_id"]]
        append_table(sheet, state_records)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            values = [len(str(cell.value or "")) for cell in column[:200]]
            width = min(max(max(values, default=10) + 2, 12), 42)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width

    workbook.save(path)


def append_table(sheet: Any, rows: list[dict[str, Any]]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    if not rows:
        return
    headers = list(rows[0])
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245B4B")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])


def main() -> None:
    with CONFIG_PATH.open(encoding="utf-8") as source:
        filter_config = json.load(source)

    records, summaries = extract(filter_config)
    csv_path = FOLDER / f"{OUTPUT_STEM}.csv"
    json_path = FOLDER / f"{OUTPUT_STEM}.json"
    summary_path = FOLDER / f"{OUTPUT_STEM}_summary.csv"
    workbook_path = FOLDER / f"{OUTPUT_STEM}.xlsx"

    write_csv(csv_path, records)
    write_csv(summary_path, summaries)
    with json_path.open("w", encoding="utf-8") as output:
        json.dump({"summary": summaries, "records": records}, output, indent=2, ensure_ascii=False)
    write_workbook(workbook_path, records, summaries)

    print(f"\nSaved {len(records)} source records")
    for summary in summaries:
        print(
            f"{summary['state']}: {summary['records_extracted']}/"
            f"{summary['target_records']} (source total {summary['available_at_source']})"
        )
    print(f"CSV: {csv_path}")
    print(f"JSON: {json_path}")
    print(f"Summary: {summary_path}")
    print(f"Excel: {workbook_path}")


if __name__ == "__main__":
    main()